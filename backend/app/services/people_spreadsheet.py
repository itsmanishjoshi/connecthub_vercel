from __future__ import annotations

import csv
import io
import os
import re
import struct
import zlib
from typing import Any, Callable

from openpyxl import load_workbook

from app.photo_store import to_data_url

HEADER_ALIASES: dict[str, list[str]] = {
    "name": [
        "name",
        "fullname",
        "attendee",
        "speakername",
        "person",
        "attendeename",
        "speakers",
        "candidate",
        "participant",
        "decisionmakers",
        "decisionmaker",
    ],
    "designation": ["designation", "title", "role", "jobtitle", "position"],
    "company": ["company", "organization", "organisation", "org", "firm", "employer"],
    "industry": [
        "industry",
        "sector",
        "primaryindustry",
        "primarysubindustry",
        "subindustry",
    ],
    "city": [
        "city",
        "contactcity",
        "officecity",
        "basecity",
        "hqcity",
        "homecity",
        "workcity",
        "employeecity",
    ],
    "location": [
        "location",
        "place",
        "office",
        "officelocation",
        "baselocation",
        "region",
        "geography",
        "geo",
        "hq",
        "headquarters",
        "countryregion",
    ],
    "linkedin_url": ["linkedin", "linkedinurl", "linkedinprofile", "linkedinid"],
    "profile_pic_url": [
        "profilepic",
        "profilepicture",
        "photo",
        "picture",
        "image",
        "avatar",
        "photourl",
        "imageurl",
    ],
    "website_url": ["website", "websiteurl", "url", "web"],
    "key_insights": ["keyinsights", "insights", "notes", "bio", "about"],
    "ice_breakers": ["icebreakers", "icebreaker", "talkingpoints", "talkingpoint"],
    "speaker": ["speaker", "isspeaker"],
    "competitor": ["competitor"],
    "priority": ["priority", "dealpriority", "contactpriority", "tier"],
    "first": ["firstname", "first", "givenname"],
    "last": ["lastname", "last", "surname", "familyname"],
}


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def map_header(value: Any) -> str | None:
    key = normalize_header(value)
    if not key:
        return None
    for field, aliases in HEADER_ALIASES.items():
        if key in aliases:
            return field
    return None


def _looks_like_yes_no(value: Any) -> bool:
    return bool(re.match(r"^(yes|no|y|n|true|false|1|0|speaker)$", str(value or "").strip(), re.I))


def _looks_like_person_name(value: Any) -> bool:
    text = str(value or "").strip()
    if not text or len(text) < 3 or len(text) > 80:
        return False
    if re.search(r"^https?://", text, re.I) or "@" in text:
        return False
    if _looks_like_yes_no(text):
        return False
    if re.fullmatch(r"\d+", text):
        return False
    return bool(re.search(r"[a-zA-Z]{2,}", text))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _rows_from_csv_buffer(buffer: bytes) -> list[list[str]]:
    text = buffer.decode("utf-8")
    return [list(row) for row in csv.reader(io.StringIO(text))]


def _rows_from_worksheet(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append([_cell_text(cell) for cell in row])
    return rows


def _workbook_rows(buffer: bytes, filename: str | None = None) -> list[tuple[str, list[list[str]]]]:
    name = filename or ""
    if re.search(r"\.csv$", name, re.I):
        return [("csv", _rows_from_csv_buffer(buffer))]
    workbook = load_workbook(io.BytesIO(buffer), read_only=True, data_only=True)
    sheets: list[tuple[str, list[list[str]]]] = []
    for sheet_name in workbook.sheetnames:
        sheets.append((sheet_name, _rows_from_worksheet(workbook[sheet_name])))
    workbook.close()
    return sheets


def _speaker_column_looks_like_names(rows: list[list[str]], header_row: int, speaker_index: int) -> bool:
    samples = [
        (rows[index][speaker_index] if speaker_index < len(rows[index]) else "")
        for index in range(header_row + 1, min(header_row + 8, len(rows)))
    ]
    samples = [value for value in samples if str(value or "").strip()]
    if not samples:
        return False
    flags = sum(1 for value in samples if _looks_like_yes_no(value))
    names = sum(1 for value in samples if _looks_like_person_name(value))
    return names > flags


def _resolve_headers(rows: list[list[str]], header_row: int) -> tuple[list[str | None], bool]:
    cells = rows[header_row] if header_row < len(rows) else []
    headers = [map_header(cell) for cell in cells]
    speaker_column_is_name = False
    if "name" not in headers:
        speaker_index = next((index for index, field in enumerate(headers) if field == "speaker"), -1)
        if speaker_index >= 0 and _speaker_column_looks_like_names(rows, header_row, speaker_index):
            headers[speaker_index] = "name"
            speaker_column_is_name = True
    return headers, speaker_column_is_name


def _find_header_row(rows: list[list[str]]) -> int:
    for index in range(min(len(rows), 30)):
        headers, _ = _resolve_headers(rows, index)
        if "name" in headers:
            return index
        if "first" in headers and "last" in headers:
            return index
    return -1


def _row_to_person(
    headers: list[str | None],
    cells: list[str],
    normalize_person: Callable[[dict], dict | None],
    labels: list[str] | None = None,
) -> dict | None:
    labels = labels or []
    raw: dict[str, Any] = {"extra_data": {}}
    for index, field in enumerate(headers):
        value = cells[index] if index < len(cells) else None
        if value is None or str(value).strip() == "":
            continue
        text = str(value).strip()
        if field:
            raw[field] = text
        elif index < len(labels) and labels[index]:
            raw["extra_data"][str(labels[index]).strip()] = text
    if not raw.get("name"):
        combined = " ".join(part for part in [raw.get("first"), raw.get("last")] if part).strip()
        if combined:
            raw["name"] = combined
    if not raw.get("name"):
        guess = next(
            (str(cell or "").strip() for cell in (cells or []) if _looks_like_person_name(cell)),
            None,
        )
        if guess:
            raw["name"] = guess
    return normalize_person(raw)


def people_from_sheet_rows(
    rows: list[list[str]],
    normalize_person: Callable[[dict], dict | None],
) -> dict[str, Any]:
    header_row = _find_header_row(rows)
    if header_row < 0:
        return {"people": [], "header_row": -1, "headers": rows[0] if rows else [], "skipped": 0}
    headers, speaker_column_is_name = _resolve_headers(rows, header_row)
    labels = rows[header_row] if header_row < len(rows) else []
    people: list[dict] = []
    skipped = 0
    for index in range(header_row + 1, len(rows)):
        cells = rows[index] if index < len(rows) else []
        has_text = any(str(cell or "").strip() for cell in cells)
        if not has_text:
            continue
        person = _row_to_person(headers, cells, normalize_person, labels)
        if person:
            if speaker_column_is_name:
                person["speaker"] = True
            person["_sheet_row"] = index
            people.append(person)
        else:
            skipped += 1
    return {
        "people": people,
        "header_row": header_row,
        "headers": rows[header_row] if header_row < len(rows) else [],
        "skipped": skipped,
    }


def unzip_entries(buffer: bytes) -> dict[str, bytes]:
    files: dict[str, bytes] = {}
    if not buffer:
        return files
    eocd = -1
    start = max(0, len(buffer) - 22 - 65557)
    for index in range(len(buffer) - 22, start - 1, -1):
        if buffer[index : index + 4] == b"\x50\x4b\x05\x06":
            eocd = index
            break
    if eocd < 0:
        return files
    count = struct.unpack_from("<H", buffer, eocd + 10)[0]
    offset = struct.unpack_from("<I", buffer, eocd + 16)[0]
    for _ in range(count):
        if offset + 46 > len(buffer) or buffer[offset : offset + 4] != b"\x50\x4b\x01\x02":
            break
        method = struct.unpack_from("<H", buffer, offset + 10)[0]
        comp_size = struct.unpack_from("<I", buffer, offset + 20)[0]
        name_len = struct.unpack_from("<H", buffer, offset + 28)[0]
        extra_len = struct.unpack_from("<H", buffer, offset + 30)[0]
        comment_len = struct.unpack_from("<H", buffer, offset + 32)[0]
        local_off = struct.unpack_from("<I", buffer, offset + 42)[0]
        name = buffer[offset + 46 : offset + 46 + name_len].decode("utf-8", errors="replace")
        if local_off + 30 <= len(buffer):
            local_name_len = struct.unpack_from("<H", buffer, local_off + 26)[0]
            local_extra = struct.unpack_from("<H", buffer, local_off + 28)[0]
            data_start = local_off + 30 + local_name_len + local_extra
            compressed = buffer[data_start : data_start + comp_size]
            try:
                if method == 0:
                    data = compressed
                elif method == 8:
                    data = zlib.decompress(compressed, -zlib.MAX_WBITS)
                else:
                    data = None
                if data is not None:
                    files[name.replace("\\", "/")] = data
            except Exception:
                pass
        offset += 46 + name_len + extra_len + comment_len
    return files


def _media_path_from_rel(target: str) -> str:
    normalized = str(target or "").replace("\\", "/")
    if normalized.startswith("xl/"):
        return normalized
    if normalized.startswith("../"):
        stripped = re.sub(r"^(\.\./)+", "", normalized)
        return f"xl/{stripped}"
    return f"xl/media/{os.path.basename(normalized)}"


def _extract_rich_cell_photos(entries: dict[str, bytes]) -> dict[int, bytes]:
    photos: dict[int, bytes] = {}
    rels = entries.get("xl/richData/_rels/richValueRel.xml.rels", b"").decode("utf-8", errors="replace")
    rel_order = entries.get("xl/richData/richValueRel.xml", b"").decode("utf-8", errors="replace")
    values = entries.get("xl/richData/rdrichvalue.xml", b"").decode("utf-8", errors="replace")
    if not rels or not rel_order or not values:
        return photos

    rel_map: dict[str, str] = {}
    for match in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels):
        rel_map[match.group(1)] = _media_path_from_rel(match.group(2))
    rel_ids = [match.group(1) for match in re.finditer(r'r:id="([^"]+)"', rel_order)]
    image_indexes = [int(match.group(1)) for match in re.finditer(r"<rv[^>]*>\s*<v>(\d+)</v>", values)]

    for name, data in entries.items():
        if not re.search(r"xl/worksheets/sheet\d+\.xml$", name, re.I):
            continue
        sheet = data.decode("utf-8", errors="replace")
        for match in re.finditer(r'<c([^>]*vm="(\d+)"[^>]*)>', sheet):
            ref = re.search(r'r="[A-Z]+(\d+)"', match.group(1))
            if not ref:
                continue
            row = int(ref.group(1)) - 1
            rich_index = int(match.group(2)) - 1
            if rich_index < 0 or rich_index >= len(image_indexes):
                continue
            image_index = image_indexes[rich_index]
            if image_index < 0 or image_index >= len(rel_ids):
                continue
            media_path = rel_map.get(rel_ids[image_index])
            image = entries.get(media_path) if media_path else None
            if image:
                photos[row] = image
    return photos


def _extract_embedded_sheet_photos(buffer: bytes) -> dict[int, bytes]:
    photos: dict[int, bytes] = {}
    try:
        entries = unzip_entries(buffer)
    except Exception:
        return photos
    photos.update(_extract_rich_cell_photos(entries))
    drawing_paths = [name for name in entries if re.search(r"xl/drawings/drawing\d+\.xml$", name, re.I)]
    for drawing_path in drawing_paths:
        xml = entries.get(drawing_path, b"").decode("utf-8", errors="replace")
        rels_path = re.sub(r"drawings/([^/]+)$", r"drawings/_rels/\1.rels", drawing_path, flags=re.I)
        rels = entries.get(rels_path, b"").decode("utf-8", errors="replace")
        rel_map: dict[str, str] = {}
        for match in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels):
            rel_map[match.group(1)] = _media_path_from_rel(match.group(2))
        anchors = re.split(r"<xdr:(?:twoCell|oneCell)Anchor[\s>]", xml, flags=re.I)[1:]
        for anchor in anchors:
            row_match = re.search(r"<xdr:from>[\s\S]*?<xdr:row>(\d+)</xdr:row>", anchor, re.I)
            embed_match = re.search(r'r:embed="([^"]+)"', anchor, re.I)
            if not row_match or not embed_match:
                continue
            row = int(row_match.group(1))
            image = entries.get(rel_map.get(embed_match.group(1), ""))
            if image and row not in photos:
                photos[row] = image
    return photos


def _decode_xml(value: str) -> str:
    return (
        str(value or "")
        .replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
    )


def _column_index(label: str) -> int:
    index = 0
    for char in str(label or ""):
        index = index * 26 + (ord(char) - 64)
    return max(0, index - 1)


def _shared_strings_from_xml(xml: str) -> list[str]:
    return [
        "".join(_decode_xml(text.group(1)) for text in re.finditer(r"<t[^>]*>([\s\S]*?)</t>", match.group(1)))
        for match in re.finditer(r"<si>([\s\S]*?)</si>", str(xml or ""))
    ]


def _rows_from_sheet_xml(sheet_xml: str, shared: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for row_match in re.finditer(r'<row[^>]*r="(\d+)"[^>]*>([\s\S]*?)</row>', str(sheet_xml or "")):
        row_index = int(row_match.group(1)) - 1
        cells: list[str] = []
        for cell_match in re.finditer(r"<c([^>]*)>([\s\S]*?)</c>", row_match.group(2)):
            ref = re.search(r'r="([A-Z]+)(\d+)"', cell_match.group(1))
            if not ref:
                continue
            col = _column_index(ref.group(1))
            attrs = cell_match.group(1)
            inline = 't="inlineStr"' in attrs
            shared_index = 't="s"' in attrs
            text = ""
            if inline:
                text = "".join(
                    _decode_xml(item.group(1))
                    for item in re.finditer(r"<t[^>]*>([\s\S]*?)</t>", cell_match.group(2))
                )
            else:
                value = re.search(r"<v>([\s\S]*?)</v>", cell_match.group(2))
                if value:
                    text = shared[int(value.group(1))] if shared_index else _decode_xml(value.group(1))
            while len(cells) <= col:
                cells.append("")
            cells[col] = text
        while len(rows) <= row_index:
            rows.append([])
        rows[row_index] = cells
    return [row or [] for row in rows]


def _people_from_xlsx_xml(
    buffer: bytes,
    normalize_person: Callable[[dict], dict | None],
) -> dict[str, Any]:
    try:
        entries = unzip_entries(buffer)
        shared = _shared_strings_from_xml(
            entries.get("xl/sharedStrings.xml", b"").decode("utf-8", errors="replace")
        )
        collected: list[dict] = []
        seen_headers: list[str] = []
        for name, data in entries.items():
            if not re.search(r"xl/worksheets/sheet\d+\.xml$", name, re.I):
                continue
            parsed = people_from_sheet_rows(
                _rows_from_sheet_xml(data.decode("utf-8", errors="replace"), shared),
                normalize_person,
            )
            collected.extend(parsed["people"])
            headers = parsed.get("headers") or []
            if headers:
                preview = ", ".join(str(item) for item in [h for h in headers if h][:8])
                seen_headers.append(f"{os.path.basename(name)}: {preview}")
        return {"people": collected, "headers": seen_headers}
    except Exception:
        return {"people": [], "headers": []}


async def people_from_spreadsheet_buffer(
    buffer: bytes,
    options: dict | None = None,
) -> dict[str, Any]:
    from app.services.people_ingest import normalize_person

    options = options or {}
    filename = options.get("filename")
    warnings: list[str] = []
    try:
        sheet_rows = _workbook_rows(buffer, filename)
    except Exception:
        return {
            "people": [],
            "warnings": ["Could not open that Excel file. Save it as .xlsx or CSV and try again."],
        }

    collected: list[dict] = []
    seen_headers: list[str] = []
    skipped = 0
    for sheet_name, rows in sheet_rows:
        parsed = people_from_sheet_rows(rows, normalize_person)
        if len(parsed["people"]) > len(collected):
            collected = list(parsed["people"])
            skipped = parsed.get("skipped") or 0
        headers = parsed.get("headers") or []
        if headers:
            header_preview = ", ".join(str(item) for item in [h for h in headers if h][:8])
            seen_headers.append(f"{sheet_name}: {header_preview}")

    if not collected:
        from_xml = _people_from_xlsx_xml(buffer, normalize_person)
        collected.extend(from_xml["people"])
        seen_headers.extend(from_xml["headers"])

    if not collected:
        if seen_headers:
            warnings.append(f"No people rows matched. Headers found: {' | '.join(seen_headers)}")
        else:
            warnings.append(
                "The Excel file opened, but no header row with names was found. "
                "If the sheet is mostly photos, save it as CSV or paste the table."
            )
        return {"people": [], "warnings": warnings}

    if skipped:
        warnings.append(
            f"{skipped} row(s) had no readable name and were listed for review only if they had other cells."
        )

    photos = _extract_embedded_sheet_photos(buffer)
    if photos:
        attached = 0
        for person in collected:
            if person.get("profile_pic_url"):
                continue
            image = photos.get(person.get("_sheet_row"))
            if not image:
                continue
            person["profile_pic_url"] = to_data_url(image)
            attached += 1
        if attached:
            warnings.append(
                f"Imported {attached} profile {'photo' if attached == 1 else 'photos'} from the spreadsheet."
            )
    elif any(not person.get("profile_pic_url") for person in collected):
        warnings.append(
            "Profile pictures inside Excel cells are optional. Add any missing photo from the person card."
        )

    for person in collected:
        person.pop("_sheet_row", None)

    return {"people": collected, "warnings": warnings}


def is_spreadsheet_file(file: Any) -> bool:
    if isinstance(file, dict):
        name = str(file.get("originalname") or file.get("name") or "")
        mime = str(file.get("mimetype") or "")
    else:
        name = str(getattr(file, "originalname", None) or getattr(file, "name", None) or "")
        mime = str(getattr(file, "mimetype", None) or "")
    if re.search(r"\.(docx|doc|pptx|ppt)$", name, re.I) or "wordprocessing" in mime or "msword" in mime:
        return False
    return (
        "spreadsheet" in mime
        or "excel" in mime
        or bool(re.search(r"\.(xlsx|xls|ods|csv)$", name, re.I))
    )


def text_from_docx_buffer(buffer: bytes) -> str:
    xml = unzip_entries(buffer).get("word/document.xml", b"").decode("utf-8", errors="replace")
    if not xml:
        return ""
    text = (
        xml.replace("</w:p>", "\n")
        .replace("<w:tab/>", "\t")
        .replace("<w:tab/>", "\t")
    )
    text = re.sub(r"<w:tab[^/]*/>", "\t", text)
    text = re.sub(r"<w:br[^/]*/>", "\n", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = (
        text.replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
    )
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
