from __future__ import annotations

import asyncio
import csv
import io
import json
import re
from typing import Any

from pypdf import PdfReader

from app.services.ai_router import complete_text, is_ai_configured
from app.services.people_spreadsheet import (
    is_spreadsheet_file,
    normalize_header,
    people_from_spreadsheet_buffer,
    text_from_docx_buffer,
)

_CITY_EXTRA_ALIASES = {
    "city",
    "contactcity",
    "officecity",
    "basecity",
    "hqcity",
    "homecity",
    "workcity",
    "employeecity",
    "location",
    "place",
    "office",
    "officelocation",
    "baselocation",
    "region",
    "geography",
    "geo",
}
_INDUSTRY_EXTRA_ALIASES = {
    "industry",
    "sector",
    "primaryindustry",
    "primarysubindustry",
    "subindustry",
}
_PRIORITY_EXTRA_ALIASES = {"priority", "dealpriority", "contactpriority", "tier"}
_EMPTY_EXTRA_VALUES = {"", "-", "—", "n/a", "na", "none", "null", "tbd"}

KNOWN_FIELDS = [
    "name",
    "designation",
    "company",
    "industry",
    "location",
    "city",
    "profile_pic_url",
    "linkedin_url",
    "website_url",
    "key_insights",
    "ice_breakers",
    "event_association",
    "speaker",
    "competitor",
    "priority",
]

EXTRACT_PROMPT = """You extract people for a ConnectHub event.
First decide what this source is:
- roster: names with title/company/city/linkedin/speaker flags
- notes: key insights, ice breakers, talking points, bios for named people
- mix of both
Return ONLY JSON:
{"kind":"roster|notes|mix","people":[{"name":"","designation":"","company":"","industry":"","location":"","city":"","profile_pic_url":"","linkedin_url":"","website_url":"","key_insights":"","ice_breakers":"","event_association":"","speaker":false,"competitor":false,"extra_data":{}}]}
Rules:
- Never invent a person who is not named in the source.
- Copy every field exactly as written. Do not paraphrase, summarize, expand, or "fix" names, titles, companies, or locations.
- name is required and must match the source spelling exactly (including punctuation and capitalization).
- Profile write-ups belong on that person. Do not create a new person from a heading like Profile 23.
- For notes-only imports: return ONLY people explicitly named in the document. Match names exactly so they can be linked to existing cards.
- key_insights = one career/role fact per item, joined with " | ". Copy wording from the source; do not rewrite.
- ice_breakers = one question per item, joined with " | ". Copy quoted questions exactly.
- If the chunk is only notes, still return the person with those notes and leave empty roster fields empty.
- Copy city and location exactly as written. Do not rename, guess, or replace them.
- speaker true only if the source explicitly says speaker/panelist/moderator.
- profile_pic_url only if a real http(s) URL exists. Never invent a URL.
- extra_data holds other useful fields (email, session, track, country). Copy values verbatim."""


def _file_name(file: Any) -> str:
    if isinstance(file, dict):
        return str(file.get("originalname") or file.get("name") or "upload")
    return str(getattr(file, "originalname", None) or getattr(file, "name", None) or "upload")


def _file_mime(file: Any) -> str:
    if isinstance(file, dict):
        return str(file.get("mimetype") or "")
    return str(getattr(file, "mimetype", None) or "")


def _file_buffer(file: Any) -> bytes:
    if isinstance(file, dict):
        return file.get("buffer") or b""
    return getattr(file, "buffer", None) or b""


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return " | ".join(str(item).strip() for item in value if str(item).strip())
    return str(value).strip()


def as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return bool(re.match(r"^(true|yes|y|1|speaker)$", str(value or "").strip(), re.I))


def coerce_extra_data(value: Any) -> dict:
    if isinstance(value, dict):
        return value
    if isinstance(value, str) and value.strip():
        try:
            parsed = json.loads(value)
            return parsed if isinstance(parsed, dict) else {}
        except json.JSONDecodeError:
            return {}
    return {}


def extract_json(text: str | None) -> dict | None:
    if not text:
        return None
    fenced = re.search(r"```json\s*([\s\S]*?)```", text, re.I)
    raw = fenced.group(1) if fenced else text
    start = raw.find("{")
    if start == -1:
        return None
    end = raw.rfind("}")
    if end > start:
        try:
            return json.loads(raw[start : end + 1])
        except json.JSONDecodeError:
            pass
    slice_text = raw[start:]
    for suffix in ["}]}", "]}", "}"]:
        try:
            parsed = json.loads(slice_text + suffix)
            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            continue
    return None


def _as_photo_url(raw: dict) -> str | None:
    extra = raw.get("extra_data") if isinstance(raw.get("extra_data"), dict) else {}
    candidates = [
        raw.get("profile_pic_url"),
        raw.get("photo"),
        raw.get("picture"),
        raw.get("image"),
        raw.get("photo_url"),
        raw.get("image_url"),
        raw.get("avatar"),
        raw.get("profile_picture"),
        raw.get("profile_pic"),
        extra.get("photo"),
        extra.get("photo_url"),
        extra.get("image"),
        extra.get("picture"),
        extra.get("avatar"),
    ]
    for value in candidates:
        text = as_text(value)
        if not text:
            continue
        if (
            re.search(r"^https?://", text, re.I)
            or text.startswith("data:image/")
            or text.startswith("/uploads/")
            or text.startswith("/api/attendees/")
            or re.search(r"\.(jpe?g|png|webp|gif)(\?|$)", text, re.I)
        ):
            return text
    return None


def person_key(person: dict) -> str:
    return f"{as_text(person.get('name')).lower()}|{as_text(person.get('company')).lower()}"


def name_key(person: dict) -> str:
    return canonical_name_key(person)


def canonical_name_key(person: dict) -> str:
    text = as_text(person.get("name"))
    text = re.sub(r"\s*\([^)]*\)\s*", " ", text)
    text = re.sub(r"\s+", " ", text).strip().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _name_tokens(person: dict) -> list[str]:
    tokens = [part for part in canonical_name_key(person).split(" ") if part]
    if len(tokens) > 2 and len(tokens[-1]) == 1:
        return tokens[:-1]
    return tokens


def _edit_distance(left: str, right: str) -> int:
    if left == right:
        return 0
    a = str(left or "")
    b = str(right or "")
    if not a:
        return len(b)
    if not b:
        return len(a)
    row = list(range(len(b) + 1))
    for i in range(1, len(a) + 1):
        previous = i - 1
        row[0] = i
        for j in range(1, len(b) + 1):
            current = row[j]
            row[j] = previous if a[i - 1] == b[j - 1] else 1 + min(previous, row[j], row[j - 1])
            previous = current
    return row[len(b)]


def _first_names_similar(left: str, right: str) -> bool:
    if not left or not right:
        return False
    if left == right:
        return True
    shorter, longer = (left, right) if len(left) <= len(right) else (right, left)
    if len(shorter) >= 4 and longer.startswith(shorter):
        return True
    dist = _edit_distance(left, right)
    min_len = min(len(left), len(right))
    return (dist <= 2 and min_len >= 5) or (dist <= 3 and min_len >= 8)


def _last_names_compatible(left: str, right: str) -> bool:
    if not left or not right:
        return False
    if left == right:
        return True
    if len(left) == 1 and right.startswith(left):
        return True
    if len(right) == 1 and left.startswith(right):
        return True
    dist = _edit_distance(left, right)
    min_len = min(len(left), len(right))
    return (dist <= 2 and min_len >= 4) or (dist <= 3 and min_len >= 8)


def _companies_compatible(existing_company: Any, incoming_company: Any) -> bool:
    left = as_text(existing_company).lower()
    right = as_text(incoming_company).lower()
    if not left or not right or left == "—" or right == "—":
        return True
    if left == right or left in right or right in left:
        return True
    compact_left = re.sub(r"[^a-z0-9]+", "", left)
    compact_right = re.sub(r"[^a-z0-9]+", "", right)
    if compact_left == compact_right:
        return True
    dist = _edit_distance(compact_left, compact_right)
    min_len = min(len(compact_left), len(compact_right))
    return dist <= 2 and min_len >= 5


def _names_look_like_same_person(existing: dict, incoming: dict) -> bool:
    left = _name_tokens(existing)
    right = _name_tokens(incoming)
    if len(left) < 2 or len(right) < 2:
        return False
    if not _last_names_compatible(left[-1], right[-1]):
        return False
    return _first_names_similar(left[0], right[0])


def annotate_duplicate_hints(people: list[dict]) -> list[dict]:
    by_first: dict[str, list[int]] = {}
    for index, person in enumerate(people):
        tokens = _name_tokens(person)
        if not tokens:
            continue
        first = tokens[0]
        by_first.setdefault(first, []).append(index)
    for person in people:
        person["possible_duplicate"] = False
        person["duplicate_with"] = []
    groups: list[dict] = []
    for first, indexes in by_first.items():
        if len(indexes) < 2:
            continue
        groups.append(
            {
                "first": first,
                "indexes": indexes,
                "names": [people[index]["name"] for index in indexes],
            }
        )
        for index in indexes:
            people[index]["possible_duplicate"] = True
            people[index]["duplicate_with"] = [
                people[other]["name"] for other in indexes if other != index
            ]
    return groups


def merge_lines(*values: Any) -> str | None:
    parts: list[str] = []
    for value in values:
        parts.extend(re.split(r"\s*\|\s*|\n+", as_text(value)))
    cleaned = [re.sub(r"^[-*•]\s+", "", item).strip() for item in parts if item.strip()]
    unique = list(dict.fromkeys(cleaned))
    kept: list[str] = []
    shortest_first = sorted(unique, key=len)
    for item in shortest_first:
        value = item.lower()
        if any(existing.lower() == value for existing in kept):
            continue
        contained = sum(
            1 for existing in kept if len(existing) > 24 and value in existing.lower()
        )
        if contained >= 2:
            continue
        kept.append(item)
    merged = [item for item in unique if item in kept]
    return " | ".join(merged) if merged else None


def _provided_text(*values: Any) -> str | None:
    for value in values:
        text = as_text(value)
        if text:
            return text
    return None


def _parse_extra_data(value: Any) -> dict[str, str]:
    if isinstance(value, dict):
        return {str(key): as_text(item) for key, item in value.items() if key and as_text(item)}
    if isinstance(value, str) and value.strip():
        try:
            parsed = json.loads(value)
        except Exception:
            return {}
        if isinstance(parsed, dict):
            return {str(key): as_text(item) for key, item in parsed.items() if key and as_text(item)}
    return {}


def _extra_field(extra: dict[str, str], aliases: set[str]) -> str | None:
    for key, value in extra.items():
        normalized = normalize_header(key)
        if normalized not in aliases:
            continue
        text = as_text(value)
        if text and text.strip().lower() not in _EMPTY_EXTRA_VALUES:
            return text
    return None


def _normalize_priority(value: Any) -> str | None:
    text = re.sub(r"\s+", "", as_text(value).upper())
    if re.fullmatch(r"P[1-4]", text):
        return text
    return None


def promote_extra_fields(person: dict[str, Any]) -> dict[str, Any]:
    extra = _parse_extra_data(person.get("extra_data"))
    city = _provided_text(person.get("city"), extra.get("city")) or _extra_field(extra, _CITY_EXTRA_ALIASES)
    location = _provided_text(person.get("location"), person.get("place"), extra.get("location"), city) or city
    industry = (
        _provided_text(person.get("industry"), person.get("sector"), extra.get("industry"))
        or _extra_field(extra, _INDUSTRY_EXTRA_ALIASES)
    )
    priority = _normalize_priority(
        _provided_text(person.get("priority")) or _extra_field(extra, _PRIORITY_EXTRA_ALIASES)
    )
    promoted = dict(person)
    if city:
        promoted["city"] = city
    if location:
        promoted["location"] = location
    if industry:
        promoted["industry"] = industry
    if priority:
        extra["priority"] = priority
        promoted["priority"] = priority
    promoted["extra_data"] = extra
    return promoted


def find_matching_person(existing_people: list[dict], incoming: dict) -> dict | None:
    incoming_name = name_key(incoming)
    if not incoming_name:
        return None

    exact = next((row for row in existing_people if person_key(row) == person_key(incoming)), None)
    if exact:
        return exact

    same_name = [row for row in existing_people if name_key(row) == incoming_name]
    incoming_tokens = [part for part in incoming_name.split(" ") if part]
    if len(same_name) == 1 and len(incoming_tokens) >= 2:
        return same_name[0]

    incoming_company = as_text(incoming.get("company")).lower()
    same_name_compatible = [
        row for row in same_name if _companies_compatible(row.get("company"), incoming.get("company"))
    ]
    if len(same_name_compatible) == 1:
        return same_name_compatible[0]
    if incoming_company:
        company_match = next(
            (row for row in same_name_compatible if as_text(row.get("company")).lower() == incoming_company),
            None,
        )
        if company_match:
            return company_match

    incoming_parts = _name_tokens(incoming)
    if len(incoming_parts) < 2:
        return None

    fuzzy = [
        row
        for row in existing_people
        if _names_look_like_same_person(row, incoming)
        and _companies_compatible(row.get("company"), incoming.get("company"))
    ]
    if len(fuzzy) == 1:
        return fuzzy[0]

    incoming_last = incoming_parts[-1]
    by_last_name_company = [
        row
        for row in existing_people
        if len(_name_tokens(row)) >= 2
        and _last_names_compatible(_name_tokens(row)[-1], incoming_last)
        and _companies_compatible(row.get("company"), incoming.get("company"))
    ]
    if len(by_last_name_company) == 1:
        return by_last_name_company[0]
    return None


def merge_notes_onto_roster(roster: list[dict], notes: list[dict]) -> list[dict]:
    merged = [dict(person) for person in roster]
    unmatched: list[dict] = []
    for note in notes:
        match = find_matching_person(merged, note)
        if match:
            index = merged.index(match)
            merged[index] = merge_person(match, note)
        else:
            unmatched.append(note)
    return merged + unmatched


def finalize_import_people(imported: list[dict], existing_rows: list[dict] | None = None) -> list[dict]:
    existing_rows = existing_rows or []
    batch = merge_people_list(imported)
    matched_existing: dict[Any, dict] = {}
    unmatched: list[dict] = []

    for person in batch:
        existing = find_matching_person(existing_rows, person)
        if existing:
            merged = {
                **merge_person(existing, person),
                "name": existing["name"],
                "already_in_event": True,
                "existing_attendee_id": existing.get("id"),
                "ingest_action": "upsert",
            }
            prior = matched_existing.get(existing.get("id"))
            matched_existing[existing.get("id")] = merge_person(prior, merged) if prior else merged
            continue

        prior_in_batch = find_matching_person(unmatched, person)
        if prior_in_batch:
            index = unmatched.index(prior_in_batch)
            unmatched[index] = merge_person(prior_in_batch, person)
        else:
            unmatched.append(
                {
                    **person,
                    "already_in_event": False,
                    "ingest_action": "upsert",
                }
            )

    return list(matched_existing.values()) + unmatched


def merge_people_list(people: list[dict]) -> list[dict]:
    merged: list[dict] = []
    for person in people:
        match = find_matching_person(merged, person)
        if match:
            next_person = merge_person(match, person)
            merged[merged.index(match)] = next_person
        else:
            merged.append(person)
    return merged


_PLACEHOLDER_NAME = re.compile(
    r"^(?:name not visible|speaker(?:\s+name)?|attendee(?:\s+name)?|unknown|tbd|n\/a|na|none|null)$",
    re.IGNORECASE,
)


def _is_placeholder_name(name: str) -> bool:
    text = name.strip()
    if not text:
        return True
    if _PLACEHOLDER_NAME.match(text):
        return True
    return "name not visible" in text.lower()


def normalize_person(raw: Any) -> dict | None:
    if not raw or not isinstance(raw, dict):
        return None
    name = (
        re.sub(
            r"\s+",
            " ",
            re.sub(r"\s*\([^)]*\)\s*", " ", as_text(raw.get("name") or raw.get("full_name") or raw.get("attendee"))),
        )
        .strip()
    )
    if not name or _is_placeholder_name(name):
        return None
    extra: dict[str, str] = {}
    if isinstance(raw.get("extra_data"), dict):
        for key, value in raw["extra_data"].items():
            text = as_text(value)
            if key and text:
                extra[str(key)] = text
    for key, value in raw.items():
        if key in KNOWN_FIELDS or key in ("full_name", "attendee", "extra_data"):
            continue
        text = as_text(value)
        if key and text:
            extra[str(key)] = text
    country = as_text(raw.get("country") or extra.get("country") or extra.get("Country"))
    if country:
        extra["country"] = country
    return promote_extra_fields(
        {
            "name": name,
            "designation": as_text(raw.get("designation") or raw.get("title") or raw.get("role")) or None,
            "company": as_text(raw.get("company") or raw.get("organization")) or None,
            "industry": as_text(raw.get("industry") or raw.get("sector")) or None,
            "location": _provided_text(raw.get("location"), raw.get("place"), extra.get("location")),
            "city": _provided_text(raw.get("city"), extra.get("city")),
            "profile_pic_url": _as_photo_url(raw),
            "linkedin_url": as_text(raw.get("linkedin_url") or raw.get("linkedin")) or None,
            "website_url": as_text(raw.get("website_url") or raw.get("website") or raw.get("url")) or None,
            "key_insights": as_text(raw.get("key_insights") or raw.get("notes") or raw.get("bio")) or None,
            "ice_breakers": as_text(raw.get("ice_breakers") or raw.get("icebreakers") or raw.get("talking_points")) or None,
            "event_association": as_text(raw.get("event_association") or raw.get("role_at_event")) or None,
            "speaker": as_bool(raw.get("speaker")),
            "competitor": as_bool(raw.get("competitor")),
            "priority": _normalize_priority(raw.get("priority")),
            "extra_data": extra,
        }
    )


def name_from_filename(filename: str | None) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[_-]+", " ", re.sub(r"\.[^.]+$", "", str(filename or "")))).strip()


async def extract_text_from_file(file: Any) -> dict[str, Any]:
    name = _file_name(file)
    mime = _file_mime(file)
    buffer = _file_buffer(file)
    if not buffer:
        return {"text": "", "kind": "empty"}

    if mime.startswith("image/"):
        return {"text": "", "kind": "image", "filename": name, "mime": mime, "buffer": buffer}

    if "wordprocessingml" in mime or "msword" in mime or re.search(r"\.docx$", name, re.I):
        try:
            text = text_from_docx_buffer(buffer)
            return {"text": text, "kind": "docx" if text else "docx-unread", "filename": name}
        except Exception:
            return {"text": "", "kind": "docx-unread", "filename": name}

    if re.search(r"\.doc$", name, re.I):
        return {"text": "", "kind": "doc-unread", "filename": name}

    if mime == "application/pdf" or name.lower().endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(buffer))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return {"text": text or "", "kind": "pdf", "filename": name}
        except Exception:
            return {"text": "", "kind": "pdf-unread", "filename": name}

    if "spreadsheet" in mime or "excel" in mime or re.search(r"\.(xlsx|xls|ods)$", name, re.I):
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(buffer), read_only=True, data_only=True)
            sheets: list[str] = []
            for sheet_name in workbook.sheetnames:
                ws = workbook[sheet_name]
                output = io.StringIO()
                writer = csv.writer(output)
                for row in ws.iter_rows(values_only=True):
                    writer.writerow(["" if value is None else value for value in row])
                sheets.append(f"Sheet {sheet_name}:\n{output.getvalue()}")
            workbook.close()
            return {"text": "\n\n".join(sheets), "kind": "spreadsheet", "filename": name}
        except Exception:
            return {"text": "", "kind": "spreadsheet-unread", "filename": name}

    return {"text": buffer.decode("utf-8", errors="replace"), "kind": "text", "filename": name}


async def _ask_model(**kwargs: Any) -> str:
    task = kwargs.get("task", "extract")
    payload: dict[str, Any] = {
        "task": task,
        "system": kwargs.get("system"),
        "user": kwargs.get("user"),
        "images": kwargs.get("images"),
        "temperature": 0.1,
        "maxTokens": kwargs.get("maxTokens", 2500),
    }
    if task == "vision" or kwargs.get("images"):
        payload["tier"] = "heavy"
    log = kwargs.get("log")
    if log:
        payload["_log"] = log
    return await complete_text(payload)


def split_material(text: str | None, size: int = 3500) -> list[str]:
    source = str(text or "").strip()
    if not source:
        return []
    profiles = [part.strip() for part in re.split(r"(?=^Profile\s+\d+\s*[—–-])", source, flags=re.M) if part.strip()]
    if len(profiles) >= 2:
        if not re.match(r"^Profile\s+\d+", profiles[0], re.I):
            intro = profiles.pop(0)
            if profiles:
                profiles[0] = f"{intro}\n\n{profiles[0]}".strip()
        return profiles
    if len(source) <= size:
        return [source]
    parts: list[str] = []
    start = 0
    while start < len(source):
        end = min(start + size, len(source))
        if end < len(source):
            break_at = source.rfind("\n\n", start, end)
            if break_at > start + size // 3:
                end = break_at
        piece = source[start:end].strip()
        if piece:
            parts.append(piece)
        start = end
    return parts


def _detect_source_kinds(files: list[Any] | None = None, roster_count: int = 0, notes_count: int = 0) -> list[str]:
    files = files or []
    kinds: list[str] = []
    if roster_count > 0 or any(is_spreadsheet_file(file) for file in files):
        kinds.append("roster")
    if notes_count > 0 or any(re.search(r"\.docx?$", _file_name(file), re.I) for file in files):
        kinds.append("notes")
    return list(dict.fromkeys(kinds))


def _looks_like_person_name(value: Any) -> bool:
    text = as_text(value)
    if not text or len(text) < 3 or len(text) > 80:
        return False
    if re.search(r"^https?://", text, re.I) or "@" in text:
        return False
    if re.match(r"^(yes|no|true|false|speaker)$", text, re.I):
        return False
    if _is_placeholder_name(text):
        return False
    return bool(re.search(r"[a-zA-Z]{2,}", text)) and len(text.split()) >= 2


def parse_notes_document(text: str | None) -> list[dict]:
    source = str(text or "").strip()
    if not source:
        return []

    profile_pieces = [
        part.strip()
        for part in re.split(r"(?=^Profile\s+\d+\s*[—–-])", source, flags=re.I | re.M)
        if part.strip()
    ]
    if len(profile_pieces) >= 2 or re.match(r"^Profile\s+\d+", profile_pieces[0] or "", re.I):
        return [person for person in map(_parse_notes_profile, profile_pieces) if person]

    profile_dash_pieces = [
        part.strip()
        for part in re.split(r"(?=^Profile\s*[—–-]\s*[A-Z])", source, flags=re.I | re.M)
        if part.strip()
    ]
    if len(profile_dash_pieces) >= 2 or re.match(r"^Profile\s*[—–-]", profile_dash_pieces[0] or "", re.I):
        return [person for person in map(_parse_notes_profile, profile_dash_pieces) if person]

    speaker_pieces = [
        part.strip()
        for part in re.split(r"(?=^[A-Z][A-Za-z.'\-\s]{2,55}\s*[—–|]\s*\S)", source, flags=re.M)
        if part.strip()
    ]
    parsed = [person for person in map(_parse_notes_profile, speaker_pieces) if person]
    if parsed:
        return parsed

    single = _parse_notes_profile(source)
    return [single] if single else []


def _parse_notes_profile(piece: str) -> dict | None:
    lines = [
        re.sub(r"^[-*•]\s+", "", line).strip()
        for line in str(piece or "").split("\n")
        if re.sub(r"^[-*•]\s+", "", line).strip()
    ]
    if not lines:
        return None

    heading = re.sub(r"^Profile\s+\d+\s*[—–-]\s*", "", lines[0], flags=re.I)
    heading = re.sub(r"^Profile\s*[—–-]\s*", "", heading, flags=re.I).strip()
    if not heading or re.match(r"^name not visible$", heading, re.I):
        return None

    pipe_match = re.match(r"^(.+?)\s*[—–|]\s*(.+)$", heading)
    name = as_text(pipe_match.group(1) if pipe_match else heading)
    company = as_text(pipe_match.group(2) if pipe_match else "")
    if not _looks_like_person_name(name):
        return None

    ice_index = next(
        (
            index
            for index, line in enumerate(lines)
            if index > 0
            and re.search(r"ice\s*breakers?|best icebreakers?|talking points?", line, re.I)
            and len(line) < 80
        ),
        -1,
    )
    insight_lines = (
        lines[1:ice_index] if ice_index >= 0 else lines[1:]
    )
    insight_lines = [
        line
        for line in insight_lines
        if not re.match(r"^(career journey|key insights?|talking points?|highlights?|about|notes?|speaker)\s*:?$", line, re.I)
    ]
    quoted = [match.group(1).strip() for match in re.finditer(r'[“"]([^”"]{8,}?)[”"]', piece)]
    ice_lines = (
        [re.sub(r'^[“"]|[”"]$', "", line).strip() for line in lines[ice_index + 1 :] if line.strip()]
        if ice_index >= 0
        else []
    )

    return normalize_person(
        {
            "name": name,
            "company": company or None,
            "key_insights": "\n".join(insight_lines) or None,
            "ice_breakers": "\n".join(quoted if quoted else ice_lines) or None,
        }
    )


def _prefer_imported_value(incoming: Any, existing: Any) -> str | None:
    next_value = as_text(incoming)
    prev = as_text(existing)
    if not next_value or next_value in ("—", "-"):
        return prev or None
    return next_value


def merge_person(existing: dict, incoming: dict) -> dict:
    next_person = dict(existing)
    for field in KNOWN_FIELDS:
        if field in ("speaker", "competitor"):
            next_person[field] = bool(existing.get(field) or incoming.get(field))
            continue
        if field in ("key_insights", "ice_breakers"):
            next_person[field] = merge_lines(existing.get(field), incoming.get(field))
            continue
        next_person[field] = _prefer_imported_value(incoming.get(field), existing.get(field))
    next_person["extra_data"] = {
        **coerce_extra_data(existing.get("extra_data")),
        **coerce_extra_data(incoming.get("extra_data")),
    }
    if incoming.get("priority"):
        next_person["priority"] = incoming.get("priority")
    return promote_extra_fields(next_person)


async def extract_people(
    text: str | None = None,
    files: list[Any] | None = None,
    upload_dir: str | None = None,
    log: dict[str, Any] | None = None,
    **_: Any,
) -> dict[str, Any]:
    warnings: list[str] = []
    direct_people: list[dict] = []
    chunks: list[str] = []
    files = files or []
    if text and text.strip():
        chunks.append(text.strip())

    for file in files:
        if is_spreadsheet_file(file) and _file_buffer(file):
            parsed = await people_from_spreadsheet_buffer(
                _file_buffer(file),
                {"uploadDir": upload_dir, "filename": _file_name(file)},
            )
            warnings.extend(parsed.get("warnings") or [])
            if parsed.get("people"):
                direct_people.extend(parsed["people"])
            continue

        extracted = await extract_text_from_file(file)
        if extracted.get("kind") == "image":
            try:
                raw = await _ask_model(
                    task="vision",
                    maxTokens=1800,
                    log=log,
                    user=(
                        f"{EXTRACT_PROMPT}\nRead this image. If it is a roster, badge, slide, or list, extract people. "
                        f'If it is only a headshot, return one person using the filename "{extracted["filename"]}" '
                        "as a last-resort name."
                    ),
                    images=[{"mime": extracted["mime"], "buffer": extracted["buffer"]}],
                )
                chunks.append(raw)
            except Exception as error:
                guessed = name_from_filename(extracted.get("filename"))
                if guessed:
                    warnings.append(
                        f"Could not read image {extracted.get('filename')}. "
                        "Add the person manually or upload a roster sheet."
                    )
                else:
                    warnings.append(str(error) or f"Could not read image {extracted.get('filename')}")
            continue
        if extracted.get("kind") == "doc-unread":
            warnings.append(f"{extracted.get('filename')} is an old Word .doc file. Save it as .docx or paste the text.")
            continue
        kind = str(extracted.get("kind") or "")
        if kind.endswith("-unread"):
            warnings.append(
                f"Could not read {extracted.get('filename')}. Try Word (.docx), PDF, Excel, or pasted text."
            )
            continue
        if extracted.get("text", "").strip():
            chunks.append(extracted["text"].strip())

    material = "\n\n".join(chunks)
    notes_people = merge_people_list(parse_notes_document(material))
    if notes_people:
        warnings.append(
            f"Read talking points for {len(notes_people)} "
            f"{'person' if len(notes_people) == 1 else 'people'} from the notes document."
        )
    profile_headings = len(re.findall(r"Profile\s+\d+", material, re.I))
    has_roster = len(direct_people) > 0
    has_structured_notes = len(notes_people) > 0
    should_run_ai = material and not (has_roster and has_structured_notes) and (
        len(notes_people) < 2 or (profile_headings > len(notes_people) and is_ai_configured())
    )
    ai_people: list[dict] = []
    if should_run_ai:
        if not is_ai_configured():
            if not direct_people and not notes_people:
                err = Exception("AI is not configured. Set AZURE_OPENAI_API_KEY and AZURE_OPENAI_ENDPOINT, or upload an Excel/CSV roster.")
                err.status = 501  # type: ignore[attr-defined]
                raise err
            warnings.append("Notes were not read because AI is not configured. Roster rows from Excel were kept.")
        else:
            pieces = split_material(material)
            if len(pieces) > 1:
                warnings.append(f"Long document split into {len(pieces)} parts so AI can read all of it.")
            ai_errors: list[str] = []
            concurrency = min(4, len(pieces))

            async def process_piece(piece: str, index: int) -> list[dict]:
                content = await _ask_model(
                    task="extract",
                    maxTokens=1800,
                    log=log,
                    system=EXTRACT_PROMPT,
                    user=(
                        f"This is part {index + 1} of {len(pieces)}. Extract every named person and their notes from this part. Return one person per Profile heading.\n\n{piece}"
                        if len(pieces) > 1
                        else piece
                    ),
                )
                parsed = extract_json(content)
                if not parsed:
                    raise Exception(f"AI returned unreadable results for part {index + 1}. Try again.")
                chunk_people = [person for person in map(normalize_person, parsed.get("people") or []) if person]
                if not chunk_people:
                    raise Exception(f"AI found no named people in part {index + 1}.")
                return chunk_people

            for start in range(0, len(pieces), concurrency):
                batch = pieces[start : start + concurrency]
                results = await asyncio.gather(
                    *[process_piece(piece, start + offset) for offset, piece in enumerate(batch)],
                    return_exceptions=True,
                )
                for offset, result in enumerate(results):
                    index = start + offset
                    if isinstance(result, Exception):
                        message = str(result) or f"Part {index + 1} failed"
                        ai_errors.append(message)
                        warnings.append(f"Could not read part {index + 1} of {len(pieces)}: {message}")
                    else:
                        ai_people.extend(result)

            if not ai_people and ai_errors and not direct_people and not notes_people:
                err = Exception(ai_errors[0])
                err.status = 502  # type: ignore[attr-defined]
                raise err
            if not ai_people and ai_errors and direct_people:
                warnings.append("The notes document could not be read with AI. Your Excel roster is still ready to save.")

    supplemental = merge_people_list(notes_people + ai_people)
    if direct_people and supplemental:
        people = merge_notes_onto_roster(direct_people, supplemental)
        warnings.append(
            f"Combined roster ({len(direct_people)} rows) with notes document. "
            "Notes were attached to matching speakers; nothing was invented."
        )
        if len(people) > len(direct_people):
            warnings.append(
                f"{len(people) - len(direct_people)} note-only profile(s) did not match the roster "
                "and are listed separately for review."
            )
    else:
        people = merge_people_list(direct_people + supplemental)

    source_kinds = _detect_source_kinds(files, roster_count=len(direct_people), notes_count=len(notes_people))
    if len(source_kinds) == 2:
        warnings.insert(
            0,
            "Detected both a speaker roster and a notes document — notes are merged onto roster rows by name and company.",
        )
    elif "roster" in source_kinds:
        warnings.insert(0, "Detected a speaker roster (Excel/CSV).")
    elif "notes" in source_kinds:
        warnings.insert(0, "Detected a notes / ice-breaker document.")

    duplicate_groups = annotate_duplicate_hints(people)
    if duplicate_groups:
        sample = "; ".join(" and ".join(group["names"]) for group in duplicate_groups[:4])
        warnings.append(
            f"Some rows share a first name ({sample}). Review the list — only one card per person will be saved."
        )
    if not people:
        detail = " ".join(filter(None, warnings))
        err = Exception(
            detail
            or (
                "No people were found in that file or text."
                if material
                else "Nothing readable was provided. Add a CSV, Excel, PDF, image, or pasted list."
            )
        )
        err.status = 422  # type: ignore[attr-defined]
        raise err
    return {"people": people, "warnings": warnings, "sourceKinds": source_kinds}
